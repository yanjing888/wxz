"""从立项 Word 文档提取三实验条目，生成 Dify 知识库 Excel 导入表（每个实验一行）。"""
from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parents[1]
DOCX = Path(r"d:/yanjing/00 资料/智能体/26立项--物小智/牛顿环与空气劈尖三个实验整理.docx")
OUT = ROOT / "docs" / "kb" / "import" / "三实验教学知识库-Dify导入.xlsx"

EXPERIMENTS = [
    {
        "code": "newton_rings",
        "name": "牛顿环法测定平凸透镜曲率半径",
        "start_line": lambda line: line.startswith("实验一") and "牛顿环" in line,
        "end_line": lambda line: line.startswith("实验二"),
    },
    {
        "code": "air_wedge_thickness",
        "name": "空气劈尖干涉测量薄片厚度",
        "start_line": lambda line: line.startswith("实验二") and "空气劈尖" in line,
        "end_line": lambda line: line.startswith("实验三"),
    },
    {
        "code": "microscope_length_measurement",
        "name": "读数显微镜微小长度直接测量练习",
        "start_line": lambda line: line.startswith("实验三") and "读数显微镜" in line,
        "end_line": lambda line: line.startswith("附录"),
    },
]

HEADERS = [
    "experiment_code",
    "experiment_name",
    "content",
    "kb_type",
    "sectionKey",
    "keywords",
]

SECTION_MAP = {
    "一、实验目的、原理": ("purpose_principle", "一、实验目的与原理"),
    "一、实验目的": ("purpose_principle", "一、实验目的与原理"),
    "二、实验步骤": ("procedure", "二、实验步骤"),
    "三、本实验专项注意事项": ("caution", "三、本实验专项注意事项"),
    "三、注意事项": ("caution", "三、本实验专项注意事项"),
}

KEYWORDS = {
    "newton_rings": "牛顿环 曲率半径 暗环直径 JCD-3 空程差 半反镜 钠黄光 R=(D_m²-D_n²)/[4(m-n)λ]",
    "air_wedge_thickness": "空气劈尖 薄片厚度 条纹 棱边 L L' n' d=(L/L')(n'λ/2)",
    "microscope_length_measurement": "读数显微镜 微小长度 A A' 空程差 L=A'-A 鼓轮",
}


def load_paragraphs() -> list[str]:
    doc = Document(str(DOCX))
    lines = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
    for table in doc.tables:
        for row in table.rows:
            cells = [c.text.strip().replace("\n", " ") for c in row.cells if c.text.strip()]
            if cells:
                lines.append(" | ".join(cells))
    return lines


def belongs_to(lines: list[str], exp: dict) -> list[str]:
    chunk: list[str] = []
    in_exp = False
    for line in lines:
        if not in_exp and exp["start_line"](line):
            in_exp = True
            chunk.append(line)
            continue
        if in_exp and exp["end_line"](line):
            break
        if in_exp:
            chunk.append(line)
    return chunk


def extract_appendix(all_lines: list[str]) -> str:
    appendix: list[str] = []
    in_appendix = False
    for line in all_lines:
        if line.startswith("附录"):
            in_appendix = True
        if in_appendix:
            appendix.append(line)
    return "\n".join(appendix).strip()


def parse_sections(lines: list[str]) -> list[tuple[str, str, str]]:
    """返回 [(section_heading, section_type, body_text), ...]"""
    sections: list[tuple[str, str, str]] = []
    current_type = "purpose_principle"
    current_heading = "一、实验目的与原理"
    step_title = ""
    buffer: list[str] = []

    def flush():
        nonlocal buffer, current_type, current_heading, step_title
        text = "\n".join(buffer).strip()
        if not text:
            buffer = []
            return
        heading = current_heading
        if current_type == "step" and step_title:
            heading = f"### {step_title}"
        sections.append((heading, current_type, text))
        buffer = []

    for line in lines[1:]:  # 跳过实验标题行，正文里单独加 H1
        if line in SECTION_MAP or any(line.startswith(k) for k in SECTION_MAP):
            flush()
            key = line if line in SECTION_MAP else next(k for k in SECTION_MAP if line.startswith(k))
            current_type, current_heading = SECTION_MAP[key]
            step_title = ""
            continue
        m = re.match(r"^(\d+)\.\s*(.+)$", line)
        if m and current_type in ("procedure", "step"):
            flush()
            current_type = "step"
            step_title = line
            buffer = []
            continue
        buffer.append(line)

    flush()
    return sections


def build_experiment_content(exp: dict, chunk: list[str], appendix: str) -> str:
    parts = [f"# {exp['name']}", ""]
    for heading, stype, body in parse_sections(chunk):
        if stype == "step":
            parts.append(heading)
        else:
            parts.append(f"## {heading.lstrip('#').strip()}")
        parts.append(body)
        parts.append("")
    if appendix:
        parts.append("## 附录：示教 CCD 与通用课堂警示")
        parts.append(appendix)
    return "\n".join(parts).strip()


def build_rows(all_lines: list[str]) -> list[dict]:
    appendix = extract_appendix(all_lines)
    rows: list[dict] = []
    for exp in EXPERIMENTS:
        chunk = belongs_to(all_lines, exp)
        if not chunk:
            continue
        rows.append({
            "experiment_code": exp["code"],
            "experiment_name": exp["name"],
            "content": build_experiment_content(exp, chunk, appendix),
            "kb_type": "teaching",
            "sectionKey": exp["code"],
            "keywords": KEYWORDS.get(exp["code"], exp["name"]),
        })
    return rows


def write_excel(rows: list[dict]) -> None:
    OUT.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    ws = wb.active
    ws.title = "teaching_kb"

    header_fill = PatternFill("solid", fgColor="1E40AF")
    header_font = Font(color="FFFFFF", bold=True)
    for col, name in enumerate(HEADERS, 1):
        cell = ws.cell(row=1, column=col, value=name)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center")

    for r, row in enumerate(rows, 2):
        for c, key in enumerate(HEADERS, 1):
            ws.cell(row=r, column=c, value=row.get(key, ""))
            ws.cell(row=r, column=c).alignment = Alignment(wrap_text=True, vertical="top")

    widths = [28, 34, 90, 12, 28, 48]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A2"

    guide = wb.create_sheet("导入说明")
    guide.append(["说明", ""])
    guide.append(["行数", "每个实验 1 行，共 3 行"])
    guide.append(["content", "整篇实验教学文档（目的、步骤、注意事项、附录），导入 Dify 时作为单文档"])
    guide.append(["experiment_code", "元数据，报告助手检索时按此字段过滤"])
    guide.append(["导入建议", "Excel 每行一条文档；分段策略选「按段落」或整段不二次切分"])

    wb.save(OUT)


def main() -> None:
    if not DOCX.exists():
        raise SystemExit(f"找不到源文档: {DOCX}")
    lines = load_paragraphs()
    rows = build_rows(lines)
    write_excel(rows)
    print(f"已生成 {len(rows)} 行（每实验一行）-> {OUT}")


if __name__ == "__main__":
    main()
